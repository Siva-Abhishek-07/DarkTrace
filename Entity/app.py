from __future__ import annotations

import csv
import hashlib
import io
import json
import math
import os
import re
import secrets
import sqlite3
import threading
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from functools import lru_cache, wraps
from typing import Any

try:
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet
    REPORTLAB_AVAILABLE = True
except ImportError:
    REPORTLAB_AVAILABLE = False

try:
    from docx import Document
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

import requests
from flask import Flask, jsonify, request, send_from_directory, Response, make_response, session as flask_session, redirect

ROOT = os.path.dirname(os.path.abspath(__file__))
app = Flask(__name__, static_folder=ROOT, static_url_path="")


# ---------------------------------------------------------------------------
# Authentication + role-based access
# ---------------------------------------------------------------------------
from werkzeug.security import generate_password_hash, check_password_hash

import shutil
import tempfile

def _get_writable_path(filename: str) -> str:
    target_in_root = os.path.join(ROOT, filename)
    try:
        test_file = os.path.join(ROOT, ".write_test")
        with open(test_file, "w") as f:
            f.write("test")
        os.remove(test_file)
        return target_in_root
    except (OSError, PermissionError):
        temp_dir = tempfile.gettempdir()
        target_in_temp = os.path.join(temp_dir, filename)
        if not os.path.exists(target_in_temp) and os.path.exists(target_in_root):
            try:
                shutil.copy2(target_in_root, target_in_temp)
            except Exception:
                pass
        return target_in_temp

AUTH_DB = _get_writable_path("darktrace_users.db")
USERNAME_SUFFIX = "@darktrace.in"
ADMIN_USERNAME = "admin@darktrace.in"
# Admin credentials: username = admin  (suffix added automatically)  |  password = DarkTrace@Admin1
ADMIN_PASSWORD_HASH = "scrypt:32768:8:1$eTt8Dmx9Dcz3EDqk$ba3f30cf90b6f32dd96db8b6f9b9767269c8fcd195118e8f7a7ef0989ddaacfd42c497dca9649921f974023fd26f65ceb78d42fca2abaf87c8cc92d93bc17d53"

app.secret_key = os.environ.get("DARKTRACE_SECRET_KEY") or "darktrace-local-development-secret-change-me"
app.config.update(SESSION_COOKIE_HTTPONLY=True, SESSION_COOKIE_SAMESITE="Lax")

def _auth_db():
    conn = sqlite3.connect(AUTH_DB)
    conn.row_factory = sqlite3.Row
    conn.execute("""
        CREATE TABLE IF NOT EXISTS users (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            username TEXT NOT NULL COLLATE NOCASE UNIQUE,
            password_hash TEXT NOT NULL,
            role TEXT NOT NULL DEFAULT 'user',
            created_at TEXT NOT NULL,
            last_login TEXT
        )
    """)
    # Migrate databases created by the previous version.
    cols = {row[1] for row in conn.execute("PRAGMA table_info(users)").fetchall()}
    if "role" not in cols:
        conn.execute("ALTER TABLE users ADD COLUMN role TEXT NOT NULL DEFAULT 'user'")
    if "last_login" not in cols:
        conn.execute("ALTER TABLE users ADD COLUMN last_login TEXT")
    conn.execute("""
        CREATE TABLE IF NOT EXISTS investigations (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER NOT NULL,
            entity_ids TEXT NOT NULL,
            entity_count INTEGER NOT NULL,
            pair_count INTEGER NOT NULL,
            relationship_score REAL NOT NULL,
            risk_level TEXT NOT NULL,
            created_at TEXT NOT NULL,
            FOREIGN KEY(user_id) REFERENCES users(id)
        )
    """)
    conn.execute("""
        CREATE TABLE IF NOT EXISTS otp_codes (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER NOT NULL,
            code TEXT NOT NULL,
            created_at INTEGER NOT NULL,
            expires_at INTEGER NOT NULL,
            used INTEGER NOT NULL DEFAULT 0,
            FOREIGN KEY(user_id) REFERENCES users(id)
        )
    """)
    conn.execute(
        "INSERT OR IGNORE INTO users (username, password_hash, role, created_at) VALUES (?, ?, 'admin', ?)",
        (ADMIN_USERNAME, ADMIN_PASSWORD_HASH, time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime()))
    )
    conn.execute("UPDATE users SET role='admin' WHERE username = ? COLLATE NOCASE", (ADMIN_USERNAME,))
    conn.commit()
    return conn

def normalize_username(value: str) -> str:
    value = clean_text(value).strip().lower()
    if not value:
        raise ValueError("Username is required.")
    if "@" in value:
        if not value.endswith(USERNAME_SUFFIX):
            raise ValueError(f"Username must end with {USERNAME_SUFFIX}.")
        local = value[:-len(USERNAME_SUFFIX)]
    else:
        local = value
    local = local.strip()
    if not re.fullmatch(r"[a-z0-9][a-z0-9._-]{2,30}", local):
        raise ValueError("Use 3–31 letters, numbers, dots, underscores or hyphens.")
    return local + USERNAME_SUFFIX

def login_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        if "user_id" not in flask_session:
            if request.path.startswith("/api/"):
                return jsonify({"error": "Authentication required.", "authenticated": False}), 401
            return redirect("/signin")
        return view(*args, **kwargs)
    return wrapped

def admin_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        if "user_id" not in flask_session:
            if request.path.startswith("/api/"):
                return jsonify({"error": "Authentication required.", "authenticated": False}), 401
            return redirect("/signin")
        if flask_session.get("role") != "admin":
            if request.path.startswith("/api/"):
                return jsonify({"error": "Admin access required."}), 403
            return redirect("/")
        return view(*args, **kwargs)
    return wrapped

@app.get("/login")
def legacy_login_page():
    return redirect("/signin")

@app.get("/signin")
def signin_page():
    if "user_id" in flask_session:
        return redirect("/admin" if flask_session.get("role") == "admin" else "/")
    return send_from_directory(ROOT, "signin.html")

@app.get("/signup")
def signup_page():
    if "user_id" in flask_session:
        return redirect("/")
    return send_from_directory(ROOT, "signup.html")

@app.post("/api/auth/register")
def register():
    body = request.get_json(silent=True) or {}
    try:
        username = normalize_username(body.get("username", ""))
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 400
    if username == ADMIN_USERNAME:
        return jsonify({"error": "The admin username is reserved."}), 403
    password = str(body.get("password", ""))
    # Password policy: minimum 6 characters, with at least one lowercase,
    # one uppercase, one number, and one special character.
    if (len(password) < 6 or
        not re.search(r"[a-z]", password) or
        not re.search(r"[A-Z]", password) or
        not re.search(r"\d", password) or
        not re.search(r"[^A-Za-z0-9]", password)):
        return jsonify({"error": "Password must be at least 6 characters and contain one lowercase letter, one uppercase letter, one number, and one special character."}), 400
    conn = _auth_db()
    try:
        conn.execute(
            "INSERT INTO users (username, password_hash, role, created_at) VALUES (?, ?, 'user', ?)",
            (username, generate_password_hash(password), time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime()))
        )
        conn.commit()
    except sqlite3.IntegrityError:
        return jsonify({"error": "That username is already registered. Choose another username."}), 409
    finally:
        conn.close()
    return jsonify({"ok": True, "message": "Registration successful. You can now sign in.", "username": username}), 201

@app.post("/api/auth/login")
def login():
    body = request.get_json(silent=True) or {}
    role_hint = str(body.get("role_hint", "user")).strip().lower()

    try:
        username = normalize_username(body.get("username", ""))
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 400
    password = str(body.get("password", ""))

    # If admin role selected, enforce that only the fixed admin account is used.
    if role_hint == "admin":
        if username != ADMIN_USERNAME:
            return jsonify({"error": "Invalid admin credentials."}), 401

    conn = _auth_db()
    row = conn.execute("SELECT id, username, password_hash, role FROM users WHERE username = ? COLLATE NOCASE", (username,)).fetchone()
    if not row or not check_password_hash(row["password_hash"], password):
        conn.close()
        return jsonify({"error": "Invalid username or password."}), 401

    # Prevent a regular user from signing in via the admin role selector.
    if role_hint == "admin" and row["role"] != "admin":
        conn.close()
        return jsonify({"error": "Invalid admin credentials."}), 401

    otp_code = f"{secrets.randbelow(900000) + 100000}"
    now_ts = int(time.time())
    expires_at = now_ts + 300  # valid 5 minutes

    conn.execute("UPDATE otp_codes SET used=1 WHERE user_id=? AND used=0", (row["id"],))
    conn.execute(
        "INSERT INTO otp_codes (user_id, code, created_at, expires_at, used) VALUES (?, ?, ?, ?, 0)",
        (row["id"], otp_code, now_ts, expires_at)
    )
    conn.commit()
    conn.close()

    flask_session.clear()
    flask_session["pending_otp_user_id"] = row["id"]
    flask_session["pending_otp_username"] = row["username"]
    flask_session["pending_otp_role"] = row["role"]

    return jsonify({
        "ok": True,
        "otp_required": True,
        "username": row["username"],
        "temp_otp": otp_code,
        "message": "Temporary OTP generated. Please enter the 6-digit code to sign in."
    })

@app.post("/api/auth/verify-otp")
def verify_otp():
    pending_user_id = flask_session.get("pending_otp_user_id")
    pending_username = flask_session.get("pending_otp_username")
    pending_role = flask_session.get("pending_otp_role")
    
    body = request.get_json(silent=True) or {}
    code = str(body.get("code", "")).strip()

    if not pending_user_id or not pending_username:
        return jsonify({"error": "No pending login session. Please sign in again."}), 400

    if not code:
        return jsonify({"error": "OTP code is required."}), 400

    now_ts = int(time.time())
    conn = _auth_db()
    otp_row = conn.execute(
        "SELECT id, code, expires_at, used FROM otp_codes WHERE user_id=? AND code=? AND used=0 ORDER BY id DESC LIMIT 1",
        (pending_user_id, code)
    ).fetchone()

    if not otp_row:
        conn.close()
        return jsonify({"error": "Invalid OTP code. Please check and try again."}), 400

    if otp_row["expires_at"] < now_ts:
        conn.close()
        return jsonify({"error": "OTP code has expired. Please request a new OTP."}), 400

    conn.execute("UPDATE otp_codes SET used=1 WHERE id=?", (otp_row["id"],))
    now_str = time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime())
    conn.execute("UPDATE users SET last_login=? WHERE id=?", (now_str, pending_user_id))
    conn.commit()
    conn.close()

    flask_session.clear()
    flask_session["user_id"] = pending_user_id
    flask_session["username"] = pending_username
    flask_session["role"] = pending_role

    return jsonify({
        "ok": True,
        "username": pending_username,
        "role": pending_role,
        "redirect": "/admin" if pending_role == "admin" else "/"
    })

@app.post("/api/auth/resend-otp")
def resend_otp():
    pending_user_id = flask_session.get("pending_otp_user_id")
    pending_username = flask_session.get("pending_otp_username")
    
    if not pending_user_id or not pending_username:
        return jsonify({"error": "No pending login session. Please sign in again."}), 400

    otp_code = f"{secrets.randbelow(900000) + 100000}"
    now_ts = int(time.time())
    expires_at = now_ts + 300

    conn = _auth_db()
    conn.execute("UPDATE otp_codes SET used=1 WHERE user_id=? AND used=0", (pending_user_id,))
    conn.execute(
        "INSERT INTO otp_codes (user_id, code, created_at, expires_at, used) VALUES (?, ?, ?, ?, 0)",
        (pending_user_id, otp_code, now_ts, expires_at)
    )
    conn.commit()
    conn.close()

    return jsonify({
        "ok": True,
        "temp_otp": otp_code,
        "message": "A new temporary OTP has been generated."
    })

@app.post("/api/auth/logout")
def logout():
    flask_session.clear()
    return jsonify({"ok": True})

@app.get("/api/auth/me")
def auth_me():
    if "user_id" not in flask_session:
        return jsonify({"authenticated": False}), 401
    return jsonify({"authenticated": True, "username": flask_session["username"], "role": flask_session.get("role", "user")})

@app.get("/admin")
@admin_required
def admin_page():
    return send_from_directory(ROOT, "admin.html")

@app.get("/api/admin/users")
@admin_required
def admin_users():
    conn = _auth_db()
    rows = conn.execute("""
        SELECT u.id, u.username, u.role, u.created_at, u.last_login,
               COUNT(i.id) AS investigation_count,
               MAX(i.created_at) AS last_investigation
        FROM users u LEFT JOIN investigations i ON i.user_id=u.id
        GROUP BY u.id ORDER BY u.created_at DESC
    """).fetchall()
    conn.close()
    return jsonify({"users": [dict(r) for r in rows]})

@app.get("/api/admin/investigations")
@admin_required
def admin_investigations():
    conn = _auth_db()
    rows = conn.execute("""
        SELECT i.id, u.username, i.entity_ids, i.entity_count, i.pair_count,
               i.relationship_score, i.risk_level, i.created_at
        FROM investigations i JOIN users u ON u.id=i.user_id
        ORDER BY i.created_at DESC LIMIT 500
    """).fetchall()
    conn.close()
    out=[]
    for r in rows:
        item=dict(r)
        try: item["entity_ids"] = json.loads(item["entity_ids"])
        except Exception: pass
        out.append(item)
    return jsonify({"investigations": out})


WIKIDATA_API = "https://www.wikidata.org/w/api.php"
GDELT_API = "https://api.gdeltproject.org/api/v2/doc/doc"
MITRE_ATTACK_URL = "https://raw.githubusercontent.com/mitre-attack/attack-stix-data/master/enterprise-attack/enterprise-attack.json"
CISA_KEV_URL = "https://www.cisa.gov/sites/default/files/feeds/known_exploited_vulnerabilities.json"
HEADERS = {"User-Agent": "DarkTrace-Threat-Intelligence/2.0 (educational project)"}

http_session = requests.Session()
http_session.headers.update(HEADERS)
CACHE: dict[str, tuple[float, Any]] = {}
CACHE_LOCK = threading.Lock()
CACHE_TTL = 1800

LEDGER_FILE = _get_writable_path("audit_chain.json")
LEDGER_LOCK = threading.Lock()
MAX_ENTITIES = 30
LEDGER_VERSION = "2.1"


def _canonical_json(value: Any) -> str:
    return json.dumps(value, sort_keys=True, separators=(",", ":"), ensure_ascii=True)


def _sha256(value: str) -> str:
    return hashlib.sha256(value.encode("utf-8")).hexdigest()


def _load_chain() -> list[dict[str, Any]]:
    if not os.path.exists(LEDGER_FILE):
        return []
    try:
        with open(LEDGER_FILE, "r", encoding="utf-8") as fh:
            value = json.load(fh)
        return value if isinstance(value, list) else []
    except (OSError, json.JSONDecodeError):
        return []


def _block_material(block: dict[str, Any]) -> dict[str, Any]:
    return {k: block[k] for k in ("ledger_version", "index", "timestamp", "event", "data", "previous_hash")}


def _block_hash(block: dict[str, Any]) -> str:
    return _sha256(_canonical_json(_block_material(block)))


def _legacy_block_hash(block: dict[str, Any]) -> str:
    material = {k: block[k] for k in ("index", "timestamp", "event", "data", "previous_hash")}
    return _sha256(_canonical_json(material))


def verify_chain() -> tuple[bool, str, dict[str, Any]]:
    chain = _load_chain()
    previous = "0" * 64
    for expected_index, block in enumerate(chain):
        # Older DarkTrace ledgers used a different hash schema. They are not
        # automatically labelled as tampered: the user must explicitly
        # initialize a new v2.1 ledger after reviewing/archiving the old one.
        if block.get("ledger_version") != LEDGER_VERSION:
            legacy_hash = _legacy_block_hash(block) if all(k in block for k in ("index", "timestamp", "event", "data", "previous_hash", "hash")) else None
            if legacy_hash and block.get("hash") == legacy_hash:
                return False, "Legacy audit ledger detected. Initialize a new ledger before recording v2.1 investigations.", {
                    "status": "legacy", "invalidBlock": None, "legacy": True
                }
            return False, "Unsupported audit ledger format. Review and initialize a new ledger.", {
                "status": "legacy", "invalidBlock": expected_index, "legacy": True
            }
        if block.get("index") != expected_index:
            return False, f"Invalid block index at position {expected_index}.", {"status": "compromised", "invalidBlock": expected_index}
        if block.get("previous_hash") != previous:
            return False, f"Broken previous-hash link at block {expected_index}.", {
                "status": "compromised", "invalidBlock": expected_index, "expectedPreviousHash": previous, "actualPreviousHash": block.get("previous_hash")
            }
        expected_hash = _block_hash(block)
        if block.get("hash") != expected_hash:
            return False, f"Hash mismatch at block {expected_index}.", {
                "status": "compromised", "invalidBlock": expected_index, "expectedHash": expected_hash, "actualHash": block.get("hash")
            }
        previous = block.get("hash", previous)
    return True, "Chain verified successfully.", {"status": "verified", "invalidBlock": None, "legacy": False}


def append_audit_block(event: str, data: dict[str, Any]) -> dict[str, Any]:
    with LEDGER_LOCK:
        chain = _load_chain()
        if chain:
            ok, message, detail = verify_chain()
            if not ok:
                raise RuntimeError(message)
        # Do not create duplicate blocks when the same investigation is submitted again.
        # One unique investigation = one audit block.
        if chain:
            latest = chain[-1]
            latest_data = latest.get("data", {})
            if latest.get("event") == event and latest_data == data:
                return latest

        previous_hash = chain[-1]["hash"] if chain else "0" * 64
        block = {
            "ledger_version": LEDGER_VERSION,
            "index": len(chain),
            "timestamp": time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime()),
            "event": event,
            "data": data,
            "previous_hash": previous_hash,
        }
        block["hash"] = _block_hash(block)
        tmp_file = LEDGER_FILE + ".tmp"
        with open(tmp_file, "w", encoding="utf-8") as fh:
            json.dump(chain + [block], fh, indent=2)
        os.replace(tmp_file, LEDGER_FILE)
        return block


def ledger_status() -> dict[str, Any]:
    ok, message, detail = verify_chain()
    chain = _load_chain()
    latest = chain[-1] if chain else None
    status = detail.get("status", "verified" if ok else "compromised")
    return {
        "secure": ok, "verified": ok, "status": status, "message": message, "blocks": len(chain),
        "latestHash": latest.get("hash") if latest else None,
        "latestBlock": latest.get("index") if latest else None,
        "invalidBlock": detail.get("invalidBlock"),
        "expectedHash": detail.get("expectedHash"),
        "actualHash": detail.get("actualHash"),
        "legacy": detail.get("legacy", False),
        "ledgerVersion": LEDGER_VERSION,
        "chain": chain,
    }


def initialize_ledger() -> dict[str, Any]:
    with LEDGER_LOCK:
        old_chain = _load_chain()
        backup = None
        if old_chain:
            backup = LEDGER_FILE + ".legacy-" + time.strftime("%Y%m%d-%H%M%S", time.gmtime()) + ".json"
            with open(backup, "w", encoding="utf-8") as fh:
                json.dump(old_chain, fh, indent=2)
        tmp_file = LEDGER_FILE + ".tmp"
        with open(tmp_file, "w", encoding="utf-8") as fh:
            json.dump([], fh, indent=2)
        os.replace(tmp_file, LEDGER_FILE)
        return {"initialized": True, "ledgerVersion": LEDGER_VERSION, "previousBlocksArchived": len(old_chain), "backupFile": os.path.basename(backup) if backup else None}

def clean_text(value: str) -> str:
    return " ".join(str(value or "").strip().split())


def cached_get_json(url: str, *, params: dict | None = None, ttl: int = CACHE_TTL, timeout: int = 15) -> Any:
    key = url + "?" + _canonical_json(params or {})
    now = time.time()
    with CACHE_LOCK:
        hit = CACHE.get(key)
        if hit and now - hit[0] < ttl:
            return hit[1]
    r = http_session.get(url, params=params, timeout=timeout)
    r.raise_for_status()
    data = r.json()
    with CACHE_LOCK:
        CACHE[key] = (now, data)
    return data


def wikidata_search(query: str, limit: int = 8) -> list[dict[str, Any]]:
    data = cached_get_json(WIKIDATA_API, params={
        "action": "wbsearchentities", "search": clean_text(query), "language": "en",
        "uselang": "en", "format": "json", "limit": limit, "type": "item"
    }, ttl=300, timeout=12)
    return [{"id": i.get("id"), "label": i.get("label") or i.get("match", {}).get("text"),
             "description": i.get("description", ""), "url": f"https://www.wikidata.org/wiki/{i.get('id')}"}
            for i in data.get("search", [])]


@lru_cache(maxsize=256)
def get_entity(qid: str) -> dict[str, Any]:
    data = cached_get_json(WIKIDATA_API, params={
        "action": "wbgetentities", "ids": qid, "props": "labels|descriptions|claims|sitelinks",
        "languages": "en", "format": "json"
    }, ttl=1800, timeout=15)
    entity = data.get("entities", {}).get(qid)
    if not entity:
        raise ValueError(f"Wikidata entity not found: {qid}")
    return entity


def entity_label(entity: dict[str, Any]) -> str:
    return entity.get("labels", {}).get("en", {}).get("value") or entity.get("id", "Unknown")


def entity_description(entity: dict[str, Any]) -> str:
    return entity.get("descriptions", {}).get("en", {}).get("value", "")


def claim_targets(entity: dict[str, Any]) -> dict[str, set[str]]:
    out: dict[str, set[str]] = {}
    for prop, claims in entity.get("claims", {}).items():
        targets = set()
        for claim in claims:
            value = claim.get("mainsnak", {}).get("datavalue", {}).get("value")
            if isinstance(value, dict) and value.get("entity-type") == "item" and value.get("id"):
                targets.add(value["id"])
        if targets:
            out[prop] = targets
    return out


@lru_cache(maxsize=512)
def property_label(pid: str) -> str:
    try:
        data = cached_get_json(WIKIDATA_API, params={"action": "wbgetentities", "ids": pid, "props": "labels", "languages": "en", "format": "json"}, ttl=86400, timeout=8)
        return data.get("entities", {}).get(pid, {}).get("labels", {}).get("en", {}).get("value", pid)
    except Exception:
        return pid


def direct_relationship(a: dict, b: dict) -> list[dict]:
    ac, bc, aid, bid = claim_targets(a), claim_targets(b), a.get("id"), b.get("id")
    evidence = []
    for pid, targets in ac.items():
        if bid in targets:
            evidence.append({"kind": "Direct structured relationship", "property": pid, "propertyLabel": property_label(pid), "direction": f"{entity_label(a)} → {entity_label(b)}"})
    for pid, targets in bc.items():
        if aid in targets:
            evidence.append({"kind": "Direct structured relationship", "property": pid, "propertyLabel": property_label(pid), "direction": f"{entity_label(b)} → {entity_label(a)}"})
    return evidence


def shared_connections(a: dict, b: dict, max_items: int = 20) -> list[dict]:
    ac, bc = claim_targets(a), claim_targets(b)
    common = []
    for pid in set(ac) & set(bc):
        for qid in ac[pid] & bc[pid]:
            common.append({"kind": "Shared structured connection", "property": pid, "propertyLabel": property_label(pid), "connectionId": qid})
    return common[:max_items]


def gdelt_co_mentions(name_a: str, name_b: str) -> dict[str, Any]:
    try:
        data = cached_get_json(GDELT_API, params={"query": f'"{name_a}" "{name_b}"', "mode": "artlist", "format": "json", "maxrecords": 20, "timespan": "3months", "sort": "HybridRel"}, ttl=900, timeout=15)
        articles = data.get("articles", [])[:20]
        return {"available": True, "count": len(articles), "articles": [{"title": a.get("title"), "url": a.get("url"), "domain": a.get("domain"), "seendate": a.get("seendate")} for a in articles]}
    except Exception as exc:
        return {"available": False, "count": 0, "articles": [], "error": str(exc)}


def mitre_catalog() -> dict[str, Any]:
    try:
        data = cached_get_json(MITRE_ATTACK_URL, ttl=21600, timeout=30)
        objects = data.get("objects", [])
        groups = [o for o in objects if o.get("type") == "intrusion-set" and o.get("revoked") is not True]
        software = [o for o in objects if o.get("type") in {"malware", "tool"} and o.get("revoked") is not True]
        relationships = [o for o in objects if o.get("type") == "relationship" and o.get("revoked") is not True]
        by_id = {o.get("id"): o for o in objects}
        group_profiles = {}
        for g in groups:
            gid = g.get("id")
            names = {str(g.get("name", "")).lower()}
            names.update(str(x).lower() for x in g.get("aliases", []) if x)
            sw, tech, campaigns = set(), set(), set()
            for rel in relationships:
                if rel.get("source_ref") == gid:
                    target = by_id.get(rel.get("target_ref"), {})
                    typ = target.get("type")
                    if typ in {"malware", "tool"}:
                        sw.add(target.get("name", ""))
                    elif typ == "attack-pattern":
                        tech.add(target.get("external_references", [{}])[0].get("external_id", target.get("name", "")))
                    elif typ in {"campaign", "course-of-action"}:
                        campaigns.add(target.get("name", ""))
            group_profiles[gid] = {"name": g.get("name", ""), "aliases": sorted(names), "software": sorted(x for x in sw if x), "techniques": sorted(x for x in tech if x), "campaigns": sorted(x for x in campaigns if x), "description": g.get("description", "")}
        return {"groups": group_profiles, "groupCount": len(group_profiles), "softwareCount": len(software), "source": "MITRE ATT&CK", "timestamp": time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime())}
    except Exception as exc:
        return {"groups": {}, "groupCount": 0, "source": "MITRE ATT&CK", "available": False, "error": str(exc)}


def match_mitre_actor(label: str, catalog: dict[str, Any]) -> dict[str, Any] | None:
    norm = clean_text(label).lower()
    if not norm:
        return None
    best = None
    for gid, p in catalog.get("groups", {}).items():
        candidates = [p.get("name", "").lower(), *p.get("aliases", [])]
        if norm in candidates or any(norm == c or norm in c or c in norm for c in candidates if c):
            best = {"id": gid, **p}
            break
    return best


def cisa_context() -> dict[str, Any]:
    try:
        data = cached_get_json(CISA_KEV_URL, ttl=21600, timeout=30)
        vulns = data.get("vulnerabilities", [])
        return {"available": True, "count": len(vulns), "date": data.get("dateReleased"), "source": "CISA KEV"}
    except Exception as exc:
        return {"available": False, "count": 0, "source": "CISA KEV", "error": str(exc)}



def infer_entity_type(entity: dict[str, Any], mitre_match: dict[str, Any] | None = None) -> str:
    if mitre_match:
        return "Threat Actor"
    text = f"{entity_label(entity)} {entity_description(entity)}".lower()
    if any(x in text for x in ("malware", "ransomware", "trojan", "worm", "backdoor")):
        return "Malware / Tool"
    if any(x in text for x in ("vulnerability", "cve-", "security flaw")):
        return "Vulnerability"
    if any(x in text for x in ("campaign", "operation")):
        return "Campaign"
    if any(x in text for x in ("company", "corporation", "software company", "organization", "organisation")):
        return "Organization"
    if "person" in text or "human" in text:
        return "Person"
    return "Entity"

def score_pair(a: dict, b: dict, direct: list[dict], shared: list[dict], news: dict, mitre_a: dict | None, mitre_b: dict | None) -> tuple[int, dict, list[dict]]:
    evidence = []
    direct_points = min(35, len(direct) * 35)
    shared_points = min(20, len(shared) * 4)
    software_shared = sorted(set((mitre_a or {}).get("software", [])) & set((mitre_b or {}).get("software", [])))
    technique_shared = sorted(set((mitre_a or {}).get("techniques", [])) & set((mitre_b or {}).get("techniques", [])))
    campaign_shared = sorted(set((mitre_a or {}).get("campaigns", [])) & set((mitre_b or {}).get("campaigns", [])))
    alias_points = 0
    if mitre_a and mitre_b:
        alias_a = set((mitre_a.get("aliases") or [])); alias_b = set((mitre_b.get("aliases") or []))
        alias_points = min(5, len(alias_a & alias_b) * 5)

    mitre_software_points = min(18, len(software_shared) * 6)
    mitre_technique_points = min(12, len(technique_shared) * 2)
    campaign_points = min(5, len(campaign_shared) * 2)
    news_points = min(10, math.ceil(news.get("count", 0) / 2))
    raw = direct_points + shared_points + mitre_software_points + mitre_technique_points + campaign_points + alias_points + news_points
    score = min(100, int(round(raw)))

    def add(kind, source, detail, url, strength):
        evidence.append({"kind": kind, "source": source, "detail": detail, "url": url, "strength": strength})

    for d in direct:
        add("Direct relationship", "Wikidata", f"{d['direction']} via {d['propertyLabel']}", f"https://www.wikidata.org/wiki/{a['id']}", 35)
    for s in shared:
        add("Shared entity evidence", "Wikidata", f"Both entities share {s['propertyLabel']}", f"https://www.wikidata.org/wiki/{s['connectionId']}", 4)
    if software_shared:
        add("Shared malware/tools", "MITRE ATT&CK", ", ".join(software_shared[:8]), "https://attack.mitre.org/", mitre_software_points)
    if technique_shared:
        add("Shared techniques", "MITRE ATT&CK", ", ".join(technique_shared[:12]), "https://attack.mitre.org/", mitre_technique_points)
    if campaign_shared:
        add("Shared campaigns", "MITRE ATT&CK", ", ".join(campaign_shared[:8]), "https://attack.mitre.org/", campaign_points)
    if news.get("count"):
        add("News co-mentions", "GDELT", f"{news['count']} recent article result(s) mention both names", "https://www.gdeltproject.org/", news_points)

    sources = sum(bool(x) for x in [direct, shared, mitre_a and mitre_b and (software_shared or technique_shared or campaign_shared), news.get("count")])
    confidence = round(min(0.99, 0.42 + sources * 0.12 + min(0.16, len(evidence) * 0.015)), 2)
    if score >= 75: rtype = "Very strong multi-source correlation"
    elif score >= 50: rtype = "Strong multi-source correlation"
    elif score >= 25: rtype = "Moderate correlation"
    elif score > 0: rtype = "Weak correlation"
    else: rtype = "No strong relationship signal"
    breakdown = {"direct": direct_points, "shared": shared_points, "software": mitre_software_points, "techniques": mitre_technique_points, "campaigns": campaign_points, "aliases": alias_points, "news": news_points, "confidence": confidence, "relationshipType": rtype}
    return score, breakdown, evidence


def risk_level(score: int, confidence: float) -> tuple[str, str]:
    adjusted = score * (0.65 + confidence * 0.35)
    if adjusted >= 75: return "HIGH", "Strong project-generated correlation signal"
    if adjusted >= 50: return "ELEVATED", "Meaningful correlation requiring review"
    if adjusted >= 25: return "GUARDED", "Limited or moderate correlation"
    return "LOW", "Little source-backed relationship signal"



# ---------------------------------------------------------------------------
# PWA static files — must be served without auth at root scope
# ---------------------------------------------------------------------------
@app.get("/sw.js")
def service_worker():
    """Service worker must be served with correct MIME type and no-cache headers."""
    resp = make_response(send_from_directory(ROOT, "sw.js"))
    resp.headers["Content-Type"] = "application/javascript; charset=utf-8"
    resp.headers["Service-Worker-Allowed"] = "/"
    resp.headers["Cache-Control"] = "no-cache, no-store, must-revalidate"
    return resp


@app.get("/manifest.json")
def manifest():
    """Web App Manifest — served without auth."""
    resp = make_response(send_from_directory(ROOT, "manifest.json"))
    resp.headers["Content-Type"] = "application/manifest+json; charset=utf-8"
    resp.headers["Cache-Control"] = "no-cache"
    return resp


@app.get("/icon-192.png")
def icon_192():
    return send_from_directory(ROOT, "icon-192.png")


@app.get("/icon-512.png")
def icon_512():
    return send_from_directory(ROOT, "icon-512.png")


@app.get("/")
@login_required
def index():
    return send_from_directory(ROOT, "index.html")


@app.get("/api/health")
@login_required
def health():
    return jsonify({"ok": True, "service": "DarkTrace threat intelligence API", "version": "2.0"})


@app.get("/api/search")
@login_required
def search():
    q = clean_text(request.args.get("q", ""))
    if len(q) < 2:
        return jsonify({"results": []})
    try:
        return jsonify({"results": wikidata_search(q)})
    except requests.RequestException as exc:
        return jsonify({"error": f"Wikidata search failed: {exc}"}), 502


@app.get("/api/sources")
@login_required
def sources():
    mitre = mitre_catalog()
    cisa = cisa_context()
    return jsonify({"sources": [
        {"name": "Wikidata", "status": "online", "role": "Entity resolution and structured relationships"},
        {"name": "MITRE ATT&CK", "status": "online" if mitre.get("groups") else "degraded", "role": "Actors, malware/tools, techniques and campaigns", "groups": mitre.get("groupCount", 0)},
        {"name": "CISA KEV", "status": "online" if cisa.get("available") else "degraded", "role": "Known exploited vulnerability context", "vulnerabilities": cisa.get("count", 0)},
        {"name": "GDELT", "status": "online", "role": "Recent news co-mention signal"},
    ]})


@app.get("/api/security")
@login_required
def security():
    return jsonify(ledger_status())


@app.post("/api/security/initialize")
@login_required
def security_initialize():
    try:
        result = initialize_ledger()
        payload = ledger_status()
        payload.update(result)
        return jsonify(payload)
    except OSError as exc:
        return jsonify({"error": f"Unable to initialize ledger: {exc}"}), 500


@app.post("/api/security/verify")
@login_required
def security_verify():
    ok, message, detail = verify_chain()
    payload = ledger_status()
    payload.update({"verified": ok, "message": message, **detail})
    return jsonify(payload), (200 if ok else 409)


@app.get("/api/security/blocks")
@login_required
def security_blocks():
    return jsonify({"blocks": _load_chain()})


@app.post("/api/analyze")
@login_required
def analyze():
    body = request.get_json(silent=True) or {}
    raw_ids = body.get("entityIds") if body.get("entityIds") is not None else [body.get("entity1Id", ""), body.get("entity2Id", "")]
    if not isinstance(raw_ids, list):
        return jsonify({"error": "entityIds must be a list."}), 400
    qids = []
    for value in raw_ids:
        qid = clean_text(value)
        if qid and qid not in qids: qids.append(qid)
    if len(qids) < 2: return jsonify({"error": "Select at least 2 entities."}), 400
    if len(qids) > MAX_ENTITIES: return jsonify({"error": "DarkTrace supports up to 30 entities per investigation for performance."}), 400
    if any(not re.fullmatch(r"Q\d+", qid) for qid in qids): return jsonify({"error": "Please select valid Wikidata entities from search results."}), 400
    try:
        entities = {qid: get_entity(qid) for qid in qids}
        mitre = mitre_catalog()
        cisa = cisa_context()
        pairs = [(qids[i], qids[j]) for i in range(len(qids)) for j in range(i + 1, len(qids))]

        def analyze_pair(pair):
            aid, bid = pair; a, b = entities[aid], entities[bid]
            ma, mb = match_mitre_actor(entity_label(a), mitre), match_mitre_actor(entity_label(b), mitre)
            direct, shared = direct_relationship(a, b), shared_connections(a, b)
            news = gdelt_co_mentions(entity_label(a), entity_label(b))
            score, breakdown, evidence = score_pair(a, b, direct, shared, news, ma, mb)
            return {"entity1Id": aid, "entity2Id": bid, "relationshipScore": score, "relationshipType": breakdown["relationshipType"],
                    "confidence": breakdown["confidence"], "scoreBreakdown": breakdown, "directRelationships": direct,
                    "sharedConnections": shared, "mitre": {"entity1": ma, "entity2": mb}, "news": news, "evidence": evidence}

        pair_results = []
        with ThreadPoolExecutor(max_workers=min(8, len(pairs))) as pool:
            futures = [pool.submit(analyze_pair, p) for p in pairs]
            for f in as_completed(futures): pair_results.append(f.result())
        pair_results.sort(key=lambda x: (qids.index(x["entity1Id"]), qids.index(x["entity2Id"])))

        scores = [r["relationshipScore"] for r in pair_results]
        overall_score = round(sum(scores) / len(scores)) if scores else 0
        overall_conf = round(sum(r["confidence"] for r in pair_results) / len(pair_results), 2) if pair_results else 0
        strongest = max(pair_results, key=lambda r: r["relationshipScore"]) if pair_results else None
        risk, risk_reason = risk_level(overall_score, overall_conf)
        evidence = [e for r in pair_results for e in r["evidence"]][:160]
        try:
            audit_block = append_audit_block("MULTI_ENTITY_ANALYSIS", {"entityIds": qids, "entityCount": len(qids), "pairCount": len(pair_results), "pairScores": [{"entity1Id": r["entity1Id"], "entity2Id": r["entity2Id"], "score": r["relationshipScore"]} for r in pair_results]})
        except RuntimeError as exc:
            return jsonify({"error": str(exc), "code": "LEDGER_REQUIRES_INITIALIZATION"}), 409
        try:
            conn = _auth_db()
            conn.execute(
                "INSERT INTO investigations (user_id, entity_ids, entity_count, pair_count, relationship_score, risk_level, created_at) VALUES (?, ?, ?, ?, ?, ?, ?)",
                (flask_session["user_id"], json.dumps(qids), len(qids), len(pair_results), overall_score, risk, time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime()))
            )
            conn.commit(); conn.close()
        except Exception:
            app.logger.exception("Unable to save investigation history")
        entity_payload = [{"id": qid, "label": entity_label(entities[qid]), "description": entity_description(entities[qid]), "url": f"https://www.wikidata.org/wiki/{qid}", "mitre": match_mitre_actor(entity_label(entities[qid]), mitre), "entityType": infer_entity_type(entities[qid], match_mitre_actor(entity_label(entities[qid]), mitre))} for qid in qids]

        aggregate = {k: sum(r["scoreBreakdown"].get(k, 0) for r in pair_results) for k in ["direct", "shared", "software", "techniques", "campaigns", "aliases", "news"]}
        return jsonify({
            "entities": entity_payload, "entity1": entity_payload[0], "entity2": entity_payload[1], "relationshipScore": overall_score,
            "relationshipType": "Multi-entity evidence correlation", "confidence": overall_conf, "scoreBreakdown": {**aggregate, "confidence": overall_conf, "relationshipType": "Multi-entity evidence correlation"},
            "pairResults": pair_results, "evidence": evidence, "pairCount": len(pair_results), "structuredSignals": sum(len(r["directRelationships"]) + len(r["sharedConnections"]) for r in pair_results),
            "newsSignals": sum(r["news"]["count"] for r in pair_results), "strongestPair": strongest, "risk": {"level": risk, "reason": risk_reason, "score": overall_score, "confidence": overall_conf},
            "sources": {"mitre": mitre, "cisa": cisa}, "generatedAt": time.strftime("%Y-%m-%d %H:%M:%S UTC", time.gmtime()),
            "security": {**ledger_status(), "ledgerBlock": audit_block["index"], "auditHash": audit_block["hash"], "previousHash": audit_block["previous_hash"], "tamperEvident": True},
            "methodology": "Pairwise 0–100 correlation using direct/shared structured evidence, MITRE ATT&CK actor overlaps, recent GDELT co-mentions and source diversity. Scores are analytical indicators, not proof of identity, intent, wrongdoing or attribution.",
        })
    except requests.RequestException as exc:
        return jsonify({"error": f"External source request failed: {exc}"}), 502
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 404
    except Exception as exc:
        app.logger.exception("Analysis failed")
        return jsonify({"error": f"Analysis failed: {exc}"}), 500


def _report_data(data: dict[str, Any]) -> dict[str, Any]:
    return data.get("analysis") if isinstance(data.get("analysis"), dict) else data

def _report_text(d: dict[str, Any]) -> str:
    lines = [
        "DarkTrace Investigation Report",
        f"Generated: {time.strftime('%Y-%m-%d %H:%M UTC')}",
        "",
        f"Relationship score: {d.get('relationshipScore', '—')}%",
        f"Relationship type: {d.get('relationshipType', '—')}",
        f"Confidence: {round(float(d.get('confidence', 0))*100)}%",
        f"Risk: {(d.get('risk') or {}).get('level', '—')} — {(d.get('risk') or {}).get('reason', '—')}",
        "",
        "Entities:",
    ]
    lines += [f"- {e.get('label')} ({e.get('id')})" for e in d.get("entities", [])]
    lines += ["", "Relationships:"]
    for r in d.get("pairResults", []):
        labels = {e.get("id"): e.get("label") for e in d.get("entities", [])}
        lines.append(f"- {labels.get(r.get('entity1Id'), r.get('entity1Id'))} ↔ {labels.get(r.get('entity2Id'), r.get('entity2Id'))}: {r.get('relationshipScore')}% | {round(float(r.get('confidence',0))*100)}% confidence | {r.get('relationshipType')}")
    lines += ["", "Evidence:"]
    for e in d.get("evidence", [])[:80]:
        lines.append(f"- [{e.get('source')}] {e.get('kind')}: {e.get('detail')} — {e.get('url')}")
    lines += ["", "Methodology note:", "Relationship scores are project-generated evidence-correlation indicators. They are not proof of identity, intent, wrongdoing or attribution."]
    return "\n".join(lines)

def _html_report(d: dict[str, Any]) -> str:
    esc_html = lambda x: str(x or "").replace("&","&amp;").replace("<","&lt;").replace(">","&gt;").replace('"',"&quot;")
    labels = {e.get("id"): e.get("label") for e in d.get("entities", [])}
    rows = "".join(f"<tr><td>{esc_html(labels.get(r.get('entity1Id'),r.get('entity1Id')))}</td><td>{esc_html(labels.get(r.get('entity2Id'),r.get('entity2Id')))}</td><td>{r.get('relationshipScore')}%</td><td>{round(float(r.get('confidence',0))*100)}%</td><td>{esc_html(r.get('relationshipType'))}</td></tr>" for r in d.get("pairResults", []))
    evidence = "".join(f"<li><b>{esc_html(e.get('source'))}</b> — {esc_html(e.get('kind'))}: {esc_html(e.get('detail'))}</li>" for e in d.get("evidence", [])[:80])
    return f"""<!doctype html><html><head><meta charset="utf-8"><title>DarkTrace Investigation Report</title><style>body{{font-family:Arial,sans-serif;margin:42px;color:#172033}}h1{{color:#5b43b5}}h2{{margin-top:26px}}.metric{{display:inline-block;padding:12px;margin:4px;border:1px solid #d9e1ec;border-radius:8px}}table{{border-collapse:collapse;width:100%}}td,th{{border:1px solid #ccd5df;padding:8px;text-align:left}}small{{color:#65748a}}</style></head><body><h1>DarkTrace Investigation Report</h1><small>Generated {time.strftime('%Y-%m-%d %H:%M UTC')}</small><p><span class="metric"><b>{d.get('relationshipScore','—')}%</b><br>Relationship score</span><span class="metric"><b>{round(float(d.get('confidence',0))*100)}%</b><br>Confidence</span><span class="metric"><b>{esc_html((d.get('risk') or {}).get('level','—'))}</b><br>Risk</span></p><h2>Entities</h2><ul>{''.join(f"<li>{esc_html(e.get('label'))} ({esc_html(e.get('id'))})</li>" for e in d.get('entities',[]))}</ul><h2>Relationships</h2><table><tr><th>Entity A</th><th>Entity B</th><th>Score</th><th>Confidence</th><th>Type</th></tr>{rows}</table><h2>Evidence</h2><ul>{evidence}</ul><h2>Methodology</h2><p>Relationship scores are project-generated evidence-correlation indicators, not proof of identity, intent, wrongdoing, or attribution.</p></body></html>"""

@app.post("/api/report/html")
@login_required
def report_html():
    d = _report_data(request.get_json(silent=True) or {})
    return Response(_html_report(d), mimetype="text/html", headers={"Content-Disposition":"attachment; filename=darktrace-investigation.html"})

@app.post("/api/report/pdf")
@login_required
def report_pdf():
    if not REPORTLAB_AVAILABLE:
        return jsonify({"error":"PDF support is not installed. Run: pip install -r requirements.txt"}), 500
    d = _report_data(request.get_json(silent=True) or {})
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, rightMargin=36,leftMargin=36,topMargin=36,bottomMargin=36)
    styles = getSampleStyleSheet(); story=[Paragraph("DarkTrace Investigation Report",styles["Title"]),Spacer(1,12)]
    story.append(Paragraph(f"Relationship score: {d.get('relationshipScore','—')}% · Confidence: {round(float(d.get('confidence',0))*100)}% · Risk: {(d.get('risk') or {}).get('level','—')}",styles["BodyText"]))
    story.append(Spacer(1,12)); story.append(Paragraph("Entities",styles["Heading2"]))
    story.append(Paragraph(", ".join(str(e.get("label")) for e in d.get("entities",[])),styles["BodyText"]))
    story.append(Spacer(1,12)); story.append(Paragraph("Relationships",styles["Heading2"]))
    labels={e.get("id"):e.get("label") for e in d.get("entities",[])}
    table_data=[["Entity A","Entity B","Score","Confidence","Type"]]+[[str(labels.get(r.get("entity1Id"),r.get("entity1Id"))),str(labels.get(r.get("entity2Id"),r.get("entity2Id"))),f"{r.get('relationshipScore')}%",f"{round(float(r.get('confidence',0))*100)}%",str(r.get("relationshipType",""))] for r in d.get("pairResults",[])]
    t=Table(table_data,repeatRows=1);t.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#5b43b5")),("TEXTCOLOR",(0,0),(-1,0),colors.white),("GRID",(0,0),(-1,-1),.5,colors.grey),("FONTSIZE",(0,0),(-1,-1),7),("VALIGN",(0,0),(-1,-1),"TOP") ]));story.append(t)
    story.append(Spacer(1,12));story.append(Paragraph("Evidence",styles["Heading2"]))
    for e in d.get("evidence",[])[:80]: story.append(Paragraph(f"<b>{e.get('source')}</b> — {e.get('kind')}: {e.get('detail')}",styles["BodyText"]))
    story.append(Spacer(1,12));story.append(Paragraph("Methodology: Relationship scores are project-generated evidence-correlation indicators and are not proof of identity, intent, wrongdoing, or attribution.",styles["BodyText"]))
    doc.build(story);buf.seek(0);return Response(buf.getvalue(),mimetype="application/pdf",headers={"Content-Disposition":"attachment; filename=darktrace-investigation.pdf"})

@app.post("/api/report/word")
@login_required
def report_word():
    if not DOCX_AVAILABLE:
        return jsonify({"error":"DOCX support is not installed. Run: pip install -r requirements.txt"}), 500
    d = _report_data(request.get_json(silent=True) or {})
    doc=Document();doc.add_heading("DarkTrace Investigation Report",0)
    doc.add_paragraph(f"Relationship score: {d.get('relationshipScore','—')}% | Confidence: {round(float(d.get('confidence',0))*100)}% | Risk: {(d.get('risk') or {}).get('level','—')}")
    doc.add_heading("Entities",1)
    for e in d.get("entities",[]): doc.add_paragraph(f"{e.get('label')} ({e.get('id')})",style="List Bullet")
    doc.add_heading("Relationships",1);table=doc.add_table(rows=1, cols=5);hdr=table.rows[0].cells
    for i,x in enumerate(["Entity A","Entity B","Score","Confidence","Type"]):hdr[i].text=x
    labels={e.get("id"):e.get("label") for e in d.get("entities",[])}
    for r in d.get("pairResults",[]): 
        cells=table.add_row().cells
        vals=[labels.get(r.get("entity1Id"),r.get("entity1Id")),labels.get(r.get("entity2Id"),r.get("entity2Id")),f"{r.get('relationshipScore')}%",f"{round(float(r.get('confidence',0))*100)}%",r.get("relationshipType")]
        for i,x in enumerate(vals):cells[i].text=str(x)
    doc.add_heading("Evidence",1)
    for e in d.get("evidence",[])[:80]:doc.add_paragraph(f"[{e.get('source')}] {e.get('kind')}: {e.get('detail')}")
    doc.add_paragraph("Methodology: Relationship scores are project-generated evidence-correlation indicators and are not proof of identity, intent, wrongdoing, or attribution.")
    buf=io.BytesIO();doc.save(buf);buf.seek(0);return Response(buf.getvalue(),mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",headers={"Content-Disposition":"attachment; filename=darktrace-investigation.docx"})

@app.post("/api/report")
@login_required
def report():
    d = _report_data(request.get_json(silent=True) or {})
    return Response(_html_report(d), mimetype="text/html")


if __name__ == "__main__":
    app.run(host="127.0.0.1", port=int(os.environ.get("PORT", "5000")), debug=True)
